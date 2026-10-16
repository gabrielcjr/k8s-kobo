import io
from typing import Optional, List, Dict, Any, Tuple
from datetime import datetime, timezone
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from ..utils import sanitize_filename_part

def generate_missionarios_xlsx(
    submissions: List[Dict[str, Any]],
    cidade: Optional[str] = None,
    povoado: Optional[str] = None
) -> Tuple[io.BytesIO, str]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Missionários Renova-me"

    headers = [
        "Data", "Cidade", "Povoado", "Nome", "Idade", "Telefone", "Whatsapp",
        "Gênero", "", "Recebeu a Jesus?", "", "Endereço ou referência",
        "Observações", "GeoLocalizador", "Localizador", "Missionário(a)"
    ]

    header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="C91818", end_color="C91818", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    ws.append(headers)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for idx, s in enumerate(submissions):
        row_num = idx + 2

        d_str = s.get("data_date") or ""
        if d_str and len(d_str) >= 7 and "-" in d_str:
            parts = d_str.split("-")
            d_val = f"{parts[1]}/{parts[0]}" if len(parts) >= 2 else d_str
        else:
            d_val = d_str

        tel = (s.get("telefone") or "").strip()
        tel_digits = "".join(ch for ch in tel if ch.isdigit())
        tel_val = tel_digits if tel_digits else (tel if tel else "")

        wa_formula = (
            f'=IF(ISBLANK(F{row_num}); ""; '
            f'IF(F{row_num} = 0; "Telefone não informado"; '
            f'HYPERLINK("https://wa.me/55" & F{row_num}; "Chamar no WhatsApp")))'
        )

        lat = s.get("latitude")
        lon = s.get("longitude")
        if lat is not None and lon is not None and (lat != 0.0 or lon != 0.0):
            geo_val = f"[{lat},{lon}]"
        else:
            geo_val = "[NULL,NULL]"

        map_formula = (
            f'=IF(ISBLANK(N{row_num}); ""; '
            f'IF(N{row_num} = "[NULL,NULL]"; "Localização não capturada";'
            f'IF(N{row_num} = "[0.1,0.1]"; "Localização não capturada"; '
            f'IF(N{row_num} = "[1,1]"; "Localização não capturada"; '
            f'IF(N{row_num} = "[1,2]"; "Localização não capturada";'
            f'HYPERLINK("https://www.google.com/maps?q=" & REGEXEXTRACT(N{row_num}; "-?\\d+\\.\\d+,-?\\d+\\.\\d+"); "Ver no mapa"))))))'
        )

        sub_cid = (s.get("cidade") or "").strip()
        sub_pov = (s.get("povoado") or "").strip()

        row_cidade = (cidade.strip() if cidade and cidade.strip() else None) or sub_cid or "Cruz das Almas"
        row_povoado = (povoado.strip() if povoado and povoado.strip() else None) or sub_pov or "Centro"

        ws.append([
            d_val,
            row_cidade,
            row_povoado,
            s.get("nome") or "",
            s.get("idade") if s.get("idade") is not None else "",
            tel_val,
            wa_formula,
            s.get("genero") or "",
            "",  # Coluna I (sem família)
            s.get("recebeu_jesus") or "",  # Coluna J
            "",  # Coluna K (sem bíblia doada)
            s.get("endereco") or "",  # Coluna L
            s.get("observacoes") or "",  # Coluna M
            geo_val,  # Coluna N (GeoLocalizador)
            map_formula,  # Coluna O (Localizador)
            s.get("missionario") or ""  # Coluna P
        ])

    # Auto-adjust column width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = col[0].column_letter
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename_parts = ["Missionarios_Renovame"]
    if cidade and cidade.strip():
        safe_cidade = sanitize_filename_part(cidade.strip())
        filename_parts.append(safe_cidade)
    filename_parts.append(datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S'))
    filename = f"{'_'.join(filename_parts)}.xlsx"

    return buf, filename
