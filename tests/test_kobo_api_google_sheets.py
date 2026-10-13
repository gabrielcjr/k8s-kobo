import csv
import io
import openpyxl
import pytest
from app.database import insert_submission

@pytest.fixture(autouse=True)
def seed_test_submission():
    """Ensure at least one standard test record exists in the database."""
    insert_submission(
        uuid="test-uuid-google-sheets-1",
        form_id="a7FqboN3tiXKPo6hSc8MzJ",
        version="32 (2025-09-12 17:50:03)",
        data_date="2026-10-13",
        nome="Carlos Silva",
        idade=34,
        telefone="75999887766",
        genero="Masculino",
        recebeu_jesus="Sim",
        endereco="Rua Principal, 100",
        geopoint="-12.2570 -38.9640 120.0 5.0",
        observacoes="Interessado nos estudos bíblicos",
        missionario="Missionário Teste",
        device_id="collect:seed-test",
        submitted_by="renovamechurch",
        source="kobo_collect"
    )

def test_google_sheets_csv_endpoints_unauthenticated_access(client):
    """
    Google Sheets =IMPORTDATA(...) cannot pass custom Authorization headers.
    All CSV endpoints must be accessible unauthenticated and return HTTP 200.
    """
    endpoints = [
        "/data.csv",
        "/api/export/csv",
        "/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/data.csv",
        "/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/data.csv/",
        "/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/export-settings/esjanp2i9bM7zAYduBVtzXe/data.csv",
        "/api/v1/data/a7FqboN3tiXKPo6hSc8MzJ.csv",
    ]
    for url in endpoints:
        res = client.get(url)
        assert res.status_code == 200, f"GET {url} failed with status {res.status_code}"
        assert "text/csv" in res.headers["content-type"]
        assert res.headers.get("access-control-allow-origin") == "*"

        head_res = client.head(url)
        assert head_res.status_code == 200, f"HEAD {url} failed with status {head_res.status_code}"
        assert "text/csv" in head_res.headers["content-type"]

def test_csv_structure_raw_field_names(client):
    """
    Verify exact KoBoToolbox field structure without human labels (default API behavior).
    """
    res = client.get("/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/data.csv")
    assert res.status_code == 200
    reader = csv.reader(io.StringIO(res.text))
    rows = list(reader)
    assert len(rows) >= 2  # Header + at least 1 record
    
    header = rows[0]
    expected_columns = [
        "start", "end", "today", "deviceid", "Data", "Nome", "Idade",
        "Telefone", "Genero", "Recebeu_a_Jesus", "Endereco_ou_referencia",
        "Geolocalizador", "_Geolocalizador_latitude", "_Geolocalizador_longitude",
        "_Geolocalizador_altitude", "_Geolocalizador_precision", "Observacoes",
        "Nome_do_Mission_rio_a", "__version__", "_id", "_uuid", "_submission_time",
        "_validation_status", "_notes", "_status", "_submitted_by", "_tags", "_index"
    ]
    for col in expected_columns:
        assert col in header, f"Column {col} missing in CSV header"

    # Find row with Carlos Silva
    carlos_row = None
    for r in rows[1:]:
        if "Carlos Silva" in r:
            carlos_row = r
            break
    assert carlos_row is not None
    nome_idx = header.index("Nome")
    tel_idx = header.index("Telefone")
    lat_idx = header.index("_Geolocalizador_latitude")
    lon_idx = header.index("_Geolocalizador_longitude")
    
    assert carlos_row[nome_idx] == "Carlos Silva"
    assert carlos_row[tel_idx] == "75999887766"
    assert carlos_row[lat_idx] == "-12.257"
    assert carlos_row[lon_idx] == "-38.964"

def test_csv_structure_with_human_labels(client):
    """
    Verify ?labels=true or ?lang=pt produces human-readable Brazilian Portuguese headers.
    """
    res = client.get("/data.csv?labels=true")
    assert res.status_code == 200
    reader = csv.reader(io.StringIO(res.text))
    rows = list(reader)
    header = rows[0]
    
    expected_labels = [
        "Data", "Nome", "Idade", "Telefone", "Gênero", "Recebeu Jesus?",
        "Endereço ou referência", "Geolocalizador", "Observações",
        "Nome do Missionário(a)", "_id", "_uuid", "_origem"
    ]
    for lbl in expected_labels:
        assert lbl in header, f"Human label {lbl} missing in CSV header"

def test_kobo_json_api_v2_endpoints(client):
    """
    Verify KoBoToolbox v2 JSON endpoints return exact pagination and result schema.
    """
    json_endpoints = [
        "/data.json",
        "/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/data.json",
        "/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/data",
        "/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/data/",
        "/api/v1/data/a7FqboN3tiXKPo6hSc8MzJ.json"
    ]
    for url in json_endpoints:
        res = client.get(url)
        assert res.status_code == 200, f"Endpoint {url} failed with status {res.status_code}"
        assert res.headers.get("access-control-allow-origin") == "*"
        data = res.json()
        assert "count" in data
        assert "results" in data
        assert isinstance(data["results"], list)
        assert data["count"] >= 1

        # Check item keys
        first = data["results"][0]
        assert "_id" in first
        assert "formhub/uuid" in first
        assert "meta/instanceID" in first
        assert "_xform_id_string" in first

def test_kobo_json_api_v1_array_endpoint(client):
    """
    KoBo API v1 data endpoint returns raw array directly.
    """
    res = client.get("/api/v1/data/a7FqboN3tiXKPo6hSc8MzJ")
    assert res.status_code == 200
    assert res.headers.get("access-control-allow-origin") == "*"
    data = res.json()
    assert isinstance(data, list)
    assert len(data) >= 1

def test_kobo_api_v2_asset_metadata(client):
    """
    Verify /api/v2/assets/{id} returns asset specification and form information.
    """
    res = client.get("/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ")
    assert res.status_code == 200
    data = res.json()
    assert data.get("uid") == "a7FqboN3tiXKPo6hSc8MzJ" or "Missionários" in data.get("name", "")

def test_kobo_api_v2_export_settings(client):
    """
    Verify /api/v2/assets/{id}/export-settings returns export configurations with CSV url.
    """
    res = client.get("/api/v2/assets/a7FqboN3tiXKPo6hSc8MzJ/export-settings")
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    assert len(data) >= 1
    assert data[0]["uid"] == "esjanp2i9bM7zAYduBVtzXe"
    assert "data_url_csv" in data[0]

def test_kobo_api_v1_forms_list(client):
    """
    Verify /api/v1/forms lists the active form.
    """
    res = client.get("/api/v1/forms")
    assert res.status_code == 200
    data = res.json()
    assert isinstance(data, list)
    assert len(data) >= 1
    assert data[0]["id_string"] == "a7FqboN3tiXKPo6hSc8MzJ"

def test_geojson_export(client):
    """
    Verify GeoJSON endpoint produces valid RFC 7946 FeatureCollection for GIS mapping.
    """
    res = client.get("/api/export/geojson")
    assert res.status_code == 200
    assert res.headers.get("access-control-allow-origin") == "*"
    data = res.json()
    assert data["type"] == "FeatureCollection"
    assert "features" in data
    assert len(data["features"]) >= 1

    feature = data["features"][0]
    assert feature["type"] == "Feature"
    assert feature["geometry"]["type"] == "Point"
    coords = feature["geometry"]["coordinates"]
    assert len(coords) == 2
    # [longitude, latitude]
    assert -45 < coords[0] < -30
    assert -15 < coords[1] < -5
    assert "nome" in feature["properties"]

def test_excel_export_behavior(client, valid_auth_headers):
    """
    Verify Excel (.xlsx) export rejects invalid credentials, accepts valid auth or direct web query,
    and produces a valid XLSX workbook matching the schema.
    """
    # Bad credentials must be rejected with 401
    bad_res = client.get("/api/export/xlsx", headers={"Authorization": "Basic d3Jvbmc6Y3JlZHM="})
    assert bad_res.status_code == 401
    assert 'Basic realm="KoBoToolbox"' in bad_res.headers.get("www-authenticate", "")

    # Authorized request returns 200 and valid XLSX binary
    res = client.get("/api/export/xlsx", headers=valid_auth_headers)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    
    # Load with openpyxl to verify valid XLSX binary structure
    wb = openpyxl.load_workbook(io.BytesIO(res.content))
    assert "Missionários Renova-me" in wb.sheetnames
    ws = wb["Missionários Renova-me"]
    assert ws.max_row >= 2
    header_vals = [cell.value for cell in ws[1]]
    assert "_id" in header_vals
    assert "Nome" in header_vals
    assert "Telefone" in header_vals

    # Direct export without auth (for Excel web query) also returns 200
    direct_res = client.get("/api/export/xlsx")
    assert direct_res.status_code == 200
